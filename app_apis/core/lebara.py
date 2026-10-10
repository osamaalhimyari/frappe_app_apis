"""Lebara B2B (https://b2b.lebara.sa) as plain Python functions -- one place that holds every request this
app knows how to make to it.

Lebara has no public API. What is called here are the requests the portal's own web pages make, replayed
inside the logged-in session kept by `app_apis.core.http` (cookie jar "lebara", encrypted, survives restarts;
the OTP login is done once from App Apis > Lebara and the hourly/5-minute scheduler jobs keep it warm).

    from app_apis.core import lebara

    lebara.find_sim(msisdn="830033599574")             # one SIM, live
    for sim in lebara.iter_subscribers():              # every SIM, paged
    lebara.fetch_invoice(2026, 9)                      # the real bill, per SIM
    lebara.loss_report("2026-09")                      # what was billed vs what the ERP knows

WHAT LEBARA HOLDS (checked against the live portal; the portal's menu is the full list):

    Subscriber list      /Services/Main/Subscriber/List     the SIMs: status, ICCID, IMEI, usage, dates
    SIM detail           /SimDetail/Ajax*                   live status, profile, location of ONE SIM
    SIM actions          /SimDetail/SuspendSIM|ResumeSIM|ActivateSIM   (HTML forms)
    Transactions         /Services/M2M/SubscriberTransaction/List   Lebara's own history of every
                                                            activate / suspend / resume / deactivate
    Groups               /Services/Main/SimGroup/List
    SMS                  /Services/Main/SmsMessges/List, /SimDetail/SendSMS
    Billing invoice      POST /Main/BillingInvoice (month, year) -> invoice number; the Excel export
                         /Main/BillingInvoice/InvoiceExcelFile?InvoiceNumber= has ONE LINE PER SIM:
                         MSISDN, tariff plan, status, amount. This is the ground truth for cost.
    Not available to this account (AccessDenied): Offer, Customer, SubCustomer lists.
    Also in the menu but not read here: Reports > Monthly Usage, Events Triggers (empty), Orders, Bulk
    actions, Transactions Dashboard.

Every function here is plain Python with no role check of its own; the HTTP layer underneath
(`app_apis.core.http`) is limited to System Manager, which is what scheduler jobs run as. For scripts and the
browser use `app_apis.core.api` (whitelisted, checked).

Nothing in this module changes anything on Lebara except `sim_action` and `send_sms`, which are named so.
"""

import io
import re

import frappe
from frappe.utils import add_months, cint, flt, get_first_day, get_last_day, getdate, now_datetime

from app_apis.core import http, store

SETTINGS = "app_apis"
JAR = "lebara"
DEFAULT_BASE = "https://b2b.lebara.sa"
PAGE = 5000  # Lebara answers up to 5000 rows a request (~7 s)

SIM_STATUS = {1: "Idle", 2: "Active", 3: "Bar", 4: "Suspend", 9: "Deactivated"}
SMS_STATUS = {0: "Pending", 1: "Under Processing", 2: "Success", 3: "Failed", 5: "Exceed Limit",
              7: "Rejected By Lebara"}
# Verified from 5,000 real transactions: the result text of each code.
TRANS_TYPE = {4: "Activate SIM", 5: "Suspend SIM", 6: "Resume SIM", 7: "Deactivate SIM"}
TRANS_STATUS = {0: "Pending", 1: "Under Processing", 2: "Success", 3: "Failed", 5: "Exceed Limit",
                7: "Rejected By Lebara"}
# The invoice prints its own words for a status.
INVOICE_STATUS = {"active": "Active", "idle": "Idle", "suspend": "Suspend", "suspended": "Suspend",
                  "deactivation": "Deactivated", "deactivated": "Deactivated", "bar": "Bar"}

SIM_COLS = ["Id", "ICCID", "Msisdn", "IMSI", "SubscriberStatusM2M1", "GroupName", "SubCustomerName",
            "ActivationDate", "LastConnectionDate", "ThisMonthUsage", "LastMonthUsage", "InDataSession", "IMEI"]
# Asked for as well: what a loss dashboard needs and the first list did not carry.
SIM_COLS_EXTRA = ["IMEI_Date", "LastDayUsage", "LastWeekUsage", "RegistrationStatus", "Apn"]
SMS_COLS = ["Id", "FromMsisdn", "ToMsisdn", "Message", "MessageType", "SmsDate", "SmsStatus", "SendResult",
            "SentByUserName"]
TRANS_COLS = ["Id", "SubscriberId", "TransTypeM2M1", "Msisdn", "Iccid", "AddDate", "AddBy", "AddByName",
              "TransDate", "TransStatus", "TransResult", "BatchID"]
AJAX_CALLS = ("AjaxBssStatus", "AjaxLiveProfile", "AjaxCharts", "AjaxLocation")


class LebaraError(Exception):
	"""Lebara could not be reached or refused. The message is safe to show."""


class SessionExpired(LebaraError):
	"""The portal session is no longer logged in; an OTP login is needed (App Apis > Lebara)."""


# --------------------------------------------------------------------------
# pure helpers (no site, no network -- unit tested)
# --------------------------------------------------------------------------


def lebara_dt(value):
	"""2026-09-16T10:20:57.000 -> '2026-09-16 10:20:57' (site time as Lebara stores it), '' -> None."""
	value = str(value or "").replace("T", " ")[:19]
	return value or None


def status_text(code) -> str:
	return SIM_STATUS.get(cint(code), "")


def invoice_status(word) -> str:
	"""The invoice's word for a status, in the same vocabulary as the SIM list ('Deactivation' ->
	'Deactivated'). An unknown word is returned as printed so nothing is silently renamed."""
	raw = str(word or "").strip()
	return INVOICE_STATUS.get(raw.lower(), raw)


def digits(value) -> str:
	return "".join(c for c in str(value or "") if c.isdigit())


def list_body(skip=0, take=100, status=None, q="", columns=None, criteria=None, sort=None) -> dict:
	"""The body of a Subscriber/List request. `q` is a part of an MSISDN / ICCID / IMSI (digits only)."""
	body = {"Skip": cint(skip), "Take": max(1, min(cint(take) or 100, PAGE)), "Sort": list(sort or ["Id"]),
	        "IncludeColumns": list(columns or SIM_COLS)}
	if status not in (None, ""):
		body["EqualityFilter"] = {"SubscriberStatusM2M1": cint(status)}
	crit = criteria
	d = digits(q)
	if d:
		like = "%" + d + "%"
		crit = [[[["Msisdn"], "like", like], "or", [["ICCID"], "like", like]], "or", [["IMSI"], "like", like]]
	if crit:
		body["Criteria"] = crit
	return body


def sim_row(e: dict) -> dict:
	"""One Subscriber/List entity -> the fields of the `Lebara SIM` doctype (no ERP link yet)."""
	return {
		"subscriber_id": cint(e.get("Id")),
		"msisdn": str(e.get("Msisdn") or ""),
		"iccid": str(e.get("ICCID") or ""),
		"imsi": str(e.get("IMSI") or ""),
		"status_code": cint(e.get("SubscriberStatusM2M1")),
		"status": status_text(e.get("SubscriberStatusM2M1")),
		"group_name": str(e.get("GroupName") or ""),
		"sub_customer": str(e.get("SubCustomerName") or ""),
		"activation_date": lebara_dt(e.get("ActivationDate")),
		"last_connection": lebara_dt(e.get("LastConnectionDate")),
		"this_month_mb": flt(e.get("ThisMonthUsage")),
		"last_month_mb": flt(e.get("LastMonthUsage")),
		"last_day_mb": flt(e.get("LastDayUsage")),
		"last_week_mb": flt(e.get("LastWeekUsage")),
		"in_data_session": cint(e.get("InDataSession")),
		"imei": str(e.get("IMEI") or ""),
		"imei_date": lebara_dt(e.get("IMEI_Date")),
		"registration_status": cint(e.get("RegistrationStatus")),
		"apn": str(e.get("Apn") or ""),
	}


def match_vehicle(row: dict, erp: dict) -> dict:
	"""Which ERP vehicle uses this SIM. The ICCID is the reliable key (Customer Vehicle.sim_serial holds
	it); the IMEI Lebara's network reports is the fallback for a vehicle with no ICCID."""
	hit = erp["iccid"].get(row.get("iccid") or "")
	how = "ICCID" if hit else ""
	if not hit and row.get("imei"):
		hit = erp["imei"].get(row["imei"])
		how = "IMEI" if hit else ""
	hit = hit or {}
	return {"erp_vehicle": hit.get("erp_vehicle") or None, "erp_customer": hit.get("erp_customer") or None,
	        "erp_plate": hit.get("erp_plate") or "", "erp_imei": hit.get("erp_imei") or "", "match_by": how}


SERIAL_SCRIPT = "lebara_serial_link"
SERIAL_FLAG = "app_apis_lebara_serial"
SERIAL_FIELDS = ("iccid", "msisdn", "imsi")
SERIAL_DEFAULTS = {"match_fields": ["iccid"], "item_codes": []}


def clean_serial_rule(raw) -> dict:
	"""The serial-link script's answer, checked: only known Lebara fields, in the order given, and an
	item list of plain strings. Anything wrong keeps the default, so one typo in a hand-edited script
	cannot unlink 25,000 SIMs."""
	out = {k: list(v) for k, v in SERIAL_DEFAULTS.items()}
	if not isinstance(raw, dict):
		return out
	for key in ("match_fields", "item_codes"):
		val = raw.get(key)
		if isinstance(val, str):
			val = val.split(",")
		if isinstance(val, (list, tuple)):
			items = [str(x).strip() for x in val if str(x).strip()]
			if key == "match_fields":
				items = [x.lower() for x in items]
				if not items or any(x not in SERIAL_FIELDS for x in items):
					continue
			out[key] = items
	return out


def link_serial(row: dict, index: dict, rule: dict) -> dict:
	"""{"serial_no", "serial_item"} for one SIM row: the first of the rule's fields whose value is a
	Serial No name in `index` ({name: item_code}), else both None."""
	for field in rule.get("match_fields") or []:
		key = str(row.get(field) or "").strip()
		if key and key in index:
			return {"serial_no": key, "serial_item": index[key] or None}
	return {"serial_no": None, "serial_item": None}


def transaction_row(e: dict) -> dict:
	"""One SubscriberTransaction entity -> the fields of the `Lebara Transaction` doctype."""
	t, s = cint(e.get("TransTypeM2M1")), e.get("TransStatus")
	return {
		"transaction_id": cint(e.get("Id")),
		"subscriber_id": cint(e.get("SubscriberId")),
		"msisdn": str(e.get("Msisdn") or ""),
		"iccid": str(e.get("Iccid") or ""),
		"trans_type": t,
		"trans_type_text": TRANS_TYPE.get(t, ""),
		"trans_status": cint(s),
		"trans_status_text": TRANS_STATUS.get(cint(s), "") if s is not None else "",
		"add_date": lebara_dt(e.get("AddDate")),
		"trans_date": lebara_dt(e.get("TransDate")),
		"add_by": cint(e.get("AddBy")),
		"add_by_name": str(e.get("AddByName") or ""),
		"batch_id": str(e.get("BatchID") or ""),
		"trans_result": str(e.get("TransResult") or "")[:500],
	}


def looks_logged_out(res) -> bool:
	"""The portal answers a dead session with its login page, a redirect to it, 401/403, or Serenity's
	own JSON error -- never with an HTTP error that says so plainly."""
	if not res:
		return True
	if "/Account/Login" in str(res.get("url") or ""):
		return True
	if res.get("status") in (401, 403):
		return True
	data = res.get("json")
	if isinstance(data, dict) and (data.get("Error") or {}).get("Code") == "NotLoggedIn":
		return True
	if data is None and "Account/Login" in str(res.get("text") or "")[:20000]:
		return True
	return False


def generic_error(res) -> bool:
	data = res.get("json")
	return isinstance(data, dict) and (data.get("Error") or {}).get("Code") == "Exception"


def parse_invoice_page(text: str):
	"""{number, date, billing_month} from the page Lebara shows after the month/year form is submitted, or
	None when no invoice is shown (a month that has none yet)."""
	flat = re.sub(r"\s+", " ", re.sub(r"<script.*?</script>|<style.*?</style>|<[^>]+>", " ", text or "", flags=re.S))
	number = re.search(r"Invoice Number\s+(\d{6,})", flat)
	if not number:
		return None
	date = re.search(r"Invoice Date\s+(\d{1,2} \w{3} \d{4})", flat)
	month = re.search(r"Billing Month\s+(\w{3} \d{4})", flat)
	return {"number": number.group(1), "date": _parse_day(date.group(1)) if date else None,
	        "billing_month": month.group(1) if month else None}


def _parse_day(text):
	from datetime import datetime

	try:
		return datetime.strptime(text.strip(), "%d %b %Y").strftime("%Y-%m-%d")
	except ValueError:
		return None


_SUMMARY_KEYS = (("previous balance", "previous_balance"), ("monthly fee", "monthly_fee"),
                 ("off-bundle usage", "off_bundle_usage"), ("adjustment", "adjustment"),
                 ("other service", "other_service"), ("total payments", "total_payments"),
                 ("vat", "vat"), ("total sar", "total"))


def parse_invoice_workbook(data: bytes) -> dict:
	"""Lebara's invoice Excel -> {"summary": {...}, "lines": [{msisdn, tariff_plan, status, amount}]}.

	Columns are found by the header NAMES in the first row of the 'Numbers' sheet, not by position, so a
	re-ordered export still reads. Raises LebaraError if the file is not an invoice."""
	import openpyxl

	try:
		wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
	except Exception as e:
		raise LebaraError("The invoice file could not be opened: %s" % str(e)[:200])
	if "Numbers" not in wb.sheetnames:
		raise LebaraError("The invoice file has no 'Numbers' sheet (sheets: %s)." % ", ".join(wb.sheetnames))

	summary = {}
	if "Summary" in wb.sheetnames:
		for row in wb["Summary"].iter_rows(values_only=True):
			cells = [c for c in row if c not in (None, "")]
			if len(cells) < 2 or not isinstance(cells[0], str):
				continue
			label = cells[0].strip().lower()
			value = cells[1]
			for prefix, key in _SUMMARY_KEYS:
				if label.startswith(prefix) and isinstance(value, (int, float)):
					summary.setdefault(key, flt(value))
			if label == "number":
				summary["number"] = str(value).strip()

	lines, index = [], None
	for row in wb["Numbers"].iter_rows(values_only=True):
		if index is None:
			head = [str(c or "").strip().lower() for c in row]
			if "msisdn" in head and "amount" in head:
				index = {"msisdn": head.index("msisdn"), "amount": head.index("amount"),
				         "plan": head.index("tariff plan") if "tariff plan" in head else None,
				         "status": head.index("status") if "status" in head else None}
			continue
		msisdn = digits(row[index["msisdn"]]) if len(row) > index["msisdn"] else ""
		if not msisdn:
			continue
		amount = row[index["amount"]] if len(row) > index["amount"] else 0
		lines.append({
			"msisdn": msisdn,
			"tariff_plan": str(row[index["plan"]] or "").strip() if index["plan"] is not None else "",
			"status": invoice_status(row[index["status"]]) if index["status"] is not None else "",
			"amount": flt(amount),
		})
	if index is None:
		raise LebaraError("The invoice file's 'Numbers' sheet has no MSISDN / Amount header.")
	return {"summary": summary, "lines": lines}


def period_of(year, month) -> str:
	return "%04d-%02d" % (cint(year), cint(month))


def previous_period(today=None):
	"""(year, month) of the month before `today` -- the newest invoice that can exist."""
	d = add_months(get_first_day(getdate(today)), -1)
	return d.year, d.month


# --------------------------------------------------------------------------
# settings and session state
# --------------------------------------------------------------------------


def _setting(field):
	return frappe.db.get_single_value(SETTINGS, field)


def base() -> str:
	return str(_setting("lebara_base_url") or DEFAULT_BASE).rstrip("/")


def enabled() -> bool:
	return bool(cint(_setting("lebara_enabled")))


def set_state(status: str, error=None):
	"""Record the session state. Committed at once, because the exception that usually follows would
	otherwise roll the new state back."""
	values = {"lebara_session_status": status}
	if status == "Logged In":
		values["lebara_last_refresh"] = now_datetime()
		values["lebara_last_error"] = ""
	if error is not None:
		values["lebara_last_error"] = str(error)[:1000]
	for k, v in values.items():
		frappe.db.set_value(SETTINGS, SETTINGS, k, v, update_modified=False)
	frappe.db.commit()


def status() -> dict:
	"""Settings and session state, no HTTP."""
	blank = "0001-01-01 00:00:00"
	return {
		"enabled": cint(_setting("lebara_enabled")),
		"status": _setting("lebara_session_status"),
		"last_refresh": str(_setting("lebara_last_refresh") or "").replace(blank, ""),
		"last_error": _setting("lebara_last_error"),
		"sims_synced_at": str(_setting("lebara_sims_synced_at") or "").replace(blank, ""),
		"sims_count": cint(_setting("lebara_sims_count")),
		"sims_sync_note": _setting("lebara_sims_sync_note"),
	}


def _notify_expired(reason):
	user = _setting("lebara_notify_user")
	if not user:
		return
	frappe.get_doc({
		"doctype": "Notification Log", "for_user": user, "type": "Alert", "document_type": SETTINGS,
		"document_name": SETTINGS,
		"subject": "Lebara session expired -- request a new OTP from App Apis > Lebara",
		"email_content": str(reason or "")[:500],
	}).insert(ignore_permissions=True)


def _expired(res):
	# flip Logged In -> Expired once, so the alert goes out once
	if _setting("lebara_session_status") == "Logged In":
		where = res.get("url") if res else ""
		set_state("Expired", "Session expired at %s (%s)" % (frappe.utils.now(), where))
		_notify_expired(where)
		frappe.db.commit()
	raise SessionExpired("Lebara session expired. Request a new OTP from App Apis > Lebara.")


# --------------------------------------------------------------------------
# requests
# --------------------------------------------------------------------------


def _post_json(path, body, timeout=180):
	return http.request(jar=JAR, url=base() + path, method="POST", body=body,
	                    headers={"X-Requested-With": "XMLHttpRequest"}, csrf_cookie="CSRF-TOKEN",
	                    timeout=timeout)


def refresh_csrf():
	"""A logged-in session gets a NEW CSRF-TOKEN cookie only with a page load; every service call made
	with the pre-login token answers HTTP 500 {"Error": {"Code": "Exception"}}."""
	http.request(jar=JAR, url=base() + "/Main/Subscriber", method="GET", timeout=60)


def service(path: str, body: dict) -> dict:
	"""A Serenity JSON service call. A stale CSRF token gets one page load and one retry; a first
	"logged out" answer gets the same single re-check before the session is declared Expired, so one
	transient 401/403 does not stop every sync until somebody types a new OTP."""
	res = _post_json(path, body)
	if generic_error(res) and not looks_logged_out(res):
		refresh_csrf()
		res = _post_json(path, body)
	if res.get("error"):
		raise LebaraError("Lebara unreachable: %s" % res.get("error"))
	if looks_logged_out(res):
		refresh_csrf()
		res = _post_json(path, body)
		if res.get("error"):
			raise LebaraError("Lebara unreachable: %s" % res.get("error"))
		if looks_logged_out(res):
			_expired(res)
	data = res.get("json")
	if isinstance(data, dict) and data.get("Error"):
		raise LebaraError("Lebara: %s" % (data["Error"].get("Message") or data["Error"].get("Code")))
	if data is None:
		raise LebaraError("Lebara returned HTTP %s with no JSON" % res.get("status"))
	return data


def _paged(path: str, body: dict, page: int = PAGE):
	"""Yield every entity of a Serenity list, `page` at a time, until TotalCount or an empty page."""
	skip = cint(body.get("Skip"))
	while True:
		res = service(path, dict(body, Skip=skip, Take=page))
		rows = res.get("Entities") or []
		for r in rows:
			yield r
		skip += page
		if not rows or skip >= cint(res.get("TotalCount")):
			return


# --------------------------------------------------------------------------
# session
# --------------------------------------------------------------------------


def _credentials() -> dict:
	doc = frappe.get_doc(SETTINGS)
	user = str(doc.get("lebara_username") or "").strip()
	pwd = doc.get_password("lebara_password", raise_exception=False) if doc.get("lebara_password") else ""
	if not user or not pwd:
		raise LebaraError("Set the Lebara Username and Password in App Apis > Lebara first.")
	return {"user": user, "pwd": pwd}


def _login_post(body, otp):
	headers = {"X-Requested-With": "XMLHttpRequest"}
	if otp:
		headers["X-OTP-Request"] = "true"
	res = http.request(jar=JAR, url=base() + "/Account/Login", method="POST", body=body, headers=headers,
	                   csrf_cookie="CSRF-TOKEN")
	if res.get("error"):
		raise LebaraError("Lebara unreachable: %s" % res.get("error"))
	return res


def request_otp() -> dict:
	"""Step 1 of the login: Lebara SMSes a 4-digit code."""
	cred = _credentials()
	http.clear_jar(jar=JAR)  # a fresh session, so no half-dead cookie survives
	http.request(jar=JAR, url=base() + "/Account/Login", method="GET")
	res = _login_post({"Username": cred["user"], "Password": cred["pwd"]}, False)
	data = res.get("json") or {}
	if res.get("status") == 200 and data == {}:
		refresh_csrf()
		set_state("Logged In")
		return {"logged_in": True, "message": "Logged in without an OTP."}
	err = data.get("Error") or {}
	if err.get("Code") != "TwoFactorAuthenticationRequired":
		set_state("Login Failed", err.get("Message") or ("HTTP %s" % res.get("status")))
		raise LebaraError("Lebara login failed: %s" % (err.get("Message") or err.get("Code") or res.get("status")))
	parts = str(err.get("Arguments") or "").rsplit("|", 1)
	if len(parts) != 2:
		raise LebaraError("Lebara asked for an OTP but sent no TwoFactorGuid: %s" % err.get("Arguments"))
	frappe.db.set_value(SETTINGS, SETTINGS, "lebara_two_factor_guid", parts[1].strip(), update_modified=False)
	set_state("Waiting for OTP", "")
	return {"logged_in": False, "message": parts[0].split("<")[0].strip() or "Lebara sent a code by SMS.",
	        "seconds": 60}


def submit_otp(code) -> dict:
	"""Step 2 of the login: the code from the SMS."""
	code = digits(code)
	if not code:
		raise LebaraError("Enter the code Lebara sent by SMS.")
	guid = _setting("lebara_two_factor_guid")
	if not guid:
		raise LebaraError("No OTP request is pending. Request an OTP first.")
	cred = _credentials()
	res = _login_post({"Username": cred["user"], "Password": cred["pwd"], "TwoFactorGuid": guid,
	                   "TwoFactorCode": cint(code)}, True)
	data = res.get("json") or {}
	if res.get("status") != 200 or data.get("Error"):
		msg = (data.get("Error") or {}).get("Message") or ("HTTP %s" % res.get("status"))
		set_state("Login Failed", msg)
		raise LebaraError("Lebara rejected the code: %s. Request a new OTP and try again." % msg)
	frappe.db.set_value(SETTINGS, SETTINGS, "lebara_two_factor_guid", "", update_modified=False)
	refresh_csrf()
	set_state("Logged In")
	return {"logged_in": True, "message": data.get("PasswordExpiryMessage") or "Logged in to Lebara."}


def ping() -> dict:
	"""The keepalive: the smallest real service call. Refreshes the sliding session and proves it works."""
	service("/Services/Main/Subscriber/List", {"Take": 1, "Sort": ["Id"], "IncludeColumns": ["Id"]})
	set_state("Logged In")
	return {"logged_in": True, "at": frappe.utils.now()}


def logout() -> dict:
	http.request(jar=JAR, url=base() + "/Account/Signout", method="GET")
	http.clear_jar(jar=JAR)
	set_state("Logged Out", "")
	return {"logged_in": False}


# --------------------------------------------------------------------------
# reads: SIMs
# --------------------------------------------------------------------------


def _with_status(e: dict) -> dict:
	e["StatusText"] = status_text(e.get("SubscriberStatusM2M1"))
	return e


def subscribers_page(skip=0, take=100, status=None, q="", columns=None) -> dict:
	"""One page of the SIM list for a screen -> {"total", "rows"}."""
	res = service("/Services/Main/Subscriber/List", list_body(skip, take, status, q, columns))
	return {"total": cint(res.get("TotalCount")), "rows": [_with_status(e) for e in res.get("Entities") or []]}


def iter_subscribers(status=None, columns=None, page=PAGE):
	"""Every SIM, paged. `columns` defaults to the 13 the mirror needs plus the extra ones."""
	body = list_body(0, page, status, "", columns or (SIM_COLS + SIM_COLS_EXTRA))
	for e in _paged("/Services/Main/Subscriber/List", body, page):
		yield _with_status(e)


def find_sim(msisdn="", iccid="", imsi="", id=0, columns=None):
	"""One SIM by MSISDN, ICCID, IMSI or Lebara's own Id, or None."""
	for field, value in (("Msisdn", str(msisdn or "").strip()), ("ICCID", str(iccid or "").strip()),
	                     ("IMSI", str(imsi or "").strip()), ("Id", cint(id))):
		if value:
			body = {"Take": 1, "IncludeColumns": list(columns or (SIM_COLS + SIM_COLS_EXTRA + ["CellId", "TechnologyUsed"])),
			        "Criteria": [[field], "=", value]}
			rows = service("/Services/Main/Subscriber/List", body).get("Entities") or []
			return _with_status(rows[0]) if rows else None
	raise LebaraError("Give msisdn, iccid, imsi or id.")


def sim_detail(subscriber_id, name: str):
	"""One of the live panels of a SIM's page: AjaxBssStatus | AjaxLiveProfile | AjaxCharts | AjaxLocation."""
	if name not in AJAX_CALLS:
		raise LebaraError("Unknown SIM detail call: %s" % name)
	res = http.request(jar=JAR, url="%s/SimDetail/%s?id=%s" % (base(), name, cint(subscriber_id)), method="POST",
	                   headers={"X-Requested-With": "XMLHttpRequest"})
	if looks_logged_out(res):
		_expired(res)
	return res.get("json") or {}


def sim_full(msisdn: str):
	"""The list record plus live status, profile and location. Four requests; for one SIM on a screen."""
	sim = find_sim(msisdn=msisdn)
	if not sim:
		return None
	sim["live_status"] = sim_detail(sim["Id"], "AjaxBssStatus")
	sim["live_profile"] = sim_detail(sim["Id"], "AjaxLiveProfile")
	loc = sim_detail(sim["Id"], "AjaxLocation")
	sim["location"] = ({"lat": loc.get("latitude"), "lng": loc.get("longitude"), "map": loc.get("locLink")}
	                   if loc.get("hasMap") else None)
	return sim


def groups() -> list:
	"""The SIM groups -> [{Id, GroupName, CustomerId}]."""
	return list(_paged("/Services/Main/SimGroup/List", {"Sort": ["Id"]}))


# --------------------------------------------------------------------------
# reads: history
# --------------------------------------------------------------------------


def iter_transactions(since_id=0, page=PAGE):
	"""Lebara's own action history, oldest first, after `since_id`. One row per activate / suspend /
	resume / deactivate -- ours and the portal users' -- with who asked and how it ended."""
	body = {"Sort": ["Id"], "IncludeColumns": TRANS_COLS, "Criteria": [["Id"], ">", cint(since_id)]}
	for e in _paged("/Services/M2M/SubscriberTransaction/List", body, page):
		yield e


def sms_query(criteria=None, take=50, sort="Id DESC", msg_type=None) -> list:
	body = {"Take": cint(take) or 1, "Sort": [sort], "IncludeColumns": SMS_COLS}
	if criteria:
		body["Criteria"] = criteria
	if msg_type:
		body["EqualityFilter"] = {"MessageType": msg_type}
	rows = service("/Services/Main/SmsMessges/List", body).get("Entities") or []
	for r in rows:
		r["StatusText"] = SMS_STATUS.get(r.get("SmsStatus"), "")
		r["TypeText"] = "Receive" if r.get("MessageType") == 2 else "Send"
	return rows


def sms_history(msisdn: str) -> list:
	"""Every SMS to and from one SIM, newest first."""
	msisdn = str(msisdn or "").strip()
	if not msisdn:
		raise LebaraError("Give the msisdn.")
	crit = [[["FromMsisdn"], "=", msisdn], "or", [["ToMsisdn"], "=", msisdn]]
	body = {"Sort": ["Id DESC"], "Criteria": crit, "IncludeColumns": SMS_COLS}
	out = []
	for r in _paged("/Services/Main/SmsMessges/List", body):
		r["StatusText"] = SMS_STATUS.get(r.get("SmsStatus"), "")
		r["TypeText"] = "Receive" if r.get("MessageType") == 2 else "Send"
		out.append(r)
	return out


# --------------------------------------------------------------------------
# reads: the bill
# --------------------------------------------------------------------------


def invoice_for(year, month):
	"""{number, date, billing_month} of that month's invoice, or None when Lebara shows none."""
	res = http.form_post(jar=JAR, url=base() + "/Main/BillingInvoice",
	                     fields={"BillingMonth": str(cint(month)), "BillingYear": str(cint(year))}, timeout=120)
	if looks_logged_out({"url": res.get("url"), "status": res.get("status"), "json": None, "text": res.get("text")}):
		_expired({"url": res.get("url")})
	if res.get("error"):
		raise LebaraError("Lebara unreachable: %s" % res.get("error"))
	return parse_invoice_page(res.get("text") or "")


def invoice_excel(number: str) -> bytes:
	"""The invoice's Excel export (one line per SIM). Roughly 800 KB for 26,000 SIMs."""
	number = digits(number)
	if not number:
		raise LebaraError("Give the invoice number.")
	got = http.fetch_bytes(JAR, "%s/Main/BillingInvoice/InvoiceExcelFile?InvoiceNumber=%s" % (base(), number))
	if not got.get("ok"):
		raise LebaraError("The invoice file could not be downloaded: %s" % (got.get("error") or "HTTP %s" % got.get("status")))
	if "spreadsheet" not in (got.get("content_type") or "") and not got["content"][:2] == b"PK":
		raise LebaraError("Lebara sent something that is not a spreadsheet (%s)." % got.get("content_type"))
	return got["content"]


def fetch_invoice(year, month) -> dict:
	"""That month's invoice, read and parsed, nothing stored.

	-> {"ok", "period", "number", "date", "summary", "lines", "note"}; ok is False with a note when no
	invoice exists for the month yet."""
	period = period_of(year, month)
	head = invoice_for(year, month)
	if not head:
		return {"ok": False, "period": period, "number": "", "date": None, "summary": {}, "lines": [],
		        "note": "Lebara shows no invoice for %s yet." % period}
	parsed = parse_invoice_workbook(invoice_excel(head["number"]))
	return {"ok": True, "period": period, "number": head["number"], "date": head["date"],
	        "summary": parsed["summary"], "lines": parsed["lines"], "note": ""}


# --------------------------------------------------------------------------
# writes (the only two things here that change Lebara)
# --------------------------------------------------------------------------


def sim_action(action: str, subscriber_id) -> dict:
	"""SuspendSIM | ResumeSIM | ActivateSIM for one subscriber Id. ActivateSIM is untested on Lebara's side."""
	if action not in ("SuspendSIM", "ResumeSIM", "ActivateSIM"):
		raise LebaraError("Unknown SIM action: %s" % action)
	sub_id = cint(subscriber_id)
	if not sub_id:
		raise LebaraError("Give the subscriber id (find_sim -> Id).")
	res = http.form_post(jar=JAR, url="%s/SimDetail/%s?Id=%s" % (base(), action, sub_id), fields={"Id": sub_id})
	if looks_logged_out(res):
		_expired(res)
	messages = [a.get("message") for a in res.get("alerts") or []]
	if not res.get("ok"):
		frappe.log_error(title="Lebara %s %s failed" % (action, sub_id),
		                 message="By %s: %s" % (frappe.session.user, "; ".join(messages) or res.get("error")))
	return {"ok": bool(res.get("ok")), "message": "; ".join(messages) or res.get("error") or "",
	        "http": res.get("status")}


def send_sms(subscriber_id, msisdn: str, message: str, sender: str = "0") -> dict:
	"""An SMS to a SIM. Lebara answers {} with no message Id, so the newest row to this SIM BEFORE
	sending is returned as `before_id`: the sent SMS is the first row after it."""
	text = str(message or "")
	if not text or len(text) > 160:
		raise LebaraError("The SMS must be 1-160 characters.")
	msisdn = str(msisdn or "").strip()
	if not cint(subscriber_id) or not msisdn:
		raise LebaraError("Give id and msisdn of the destination SIM.")
	before = sms_query([["ToMsisdn"], "=", msisdn], 1, "Id DESC")
	service("/SimDetail/SendSMS", {"Id": cint(subscriber_id), "Source_MSISDN": str(sender or "0"),
	                               "Destination_MSISDN": msisdn, "Message": text})
	return {"ok": True, "before_id": cint(before[0].get("Id")) if before else 0, "message": text}


# --------------------------------------------------------------------------
# mirrors: what the dashboard reads
# --------------------------------------------------------------------------


def erp_vehicles() -> dict:
	"""Customer Vehicles keyed by SIM ICCID and by device IMEI. Ordered by `modified`, so the most
	recently edited vehicle wins a key two vehicles share."""
	by_iccid, by_imei = {}, {}
	for v in frappe.db.sql(
		"select name, customer, device_serial, sim_serial, license_plate, e_license_plate, plate_num "
		"from `tabCustomer Vehicle` where ifnull(sim_serial, '') != '' or ifnull(device_serial, '') != '' "
		"order by modified", as_dict=True):
		plate = ""
		for f in ("license_plate", "e_license_plate", "plate_num"):
			if not plate and str(v.get(f) or "").strip():
				plate = str(v.get(f)).strip()
		rec = {"erp_vehicle": v.name, "erp_customer": v.customer or "", "erp_plate": plate or v.name,
		       "erp_imei": str(v.device_serial or "").strip()}
		if str(v.sim_serial or "").strip():
			by_iccid[str(v.sim_serial).strip()] = rec
		if rec["erp_imei"]:
			by_imei[rec["erp_imei"]] = rec
	return {"iccid": by_iccid, "imei": by_imei}


def serial_rule() -> dict:
	"""The linking rule, from the lebara_serial_link Server Script when it is there and enabled, else the
	built-in one. Read once per request."""
	cached = getattr(frappe.local, "app_apis_lebara_serial_rule", None)
	if cached is not None:
		return cached
	raw = None
	try:
		row = frappe.db.get_value("Server Script", SERIAL_SCRIPT, ["disabled", "script_type"], as_dict=True)
		if row and not cint(row.disabled) and row.script_type == "API":
			answer = frappe.get_doc("Server Script", SERIAL_SCRIPT).execute_method()
			raw = (answer or {}).get(SERIAL_FLAG)
	except Exception:
		frappe.log_error(title="app_apis: lebara_serial_link script failed")
	rule = clean_serial_rule(raw)
	frappe.local.app_apis_lebara_serial_rule = rule
	return rule


def serial_index(rule: dict) -> dict:
	"""{Serial No name: Item code} for the Serial Nos the rule allows. One query, names and items only."""
	items = rule.get("item_codes") or []
	if items:
		rows = frappe.db.sql("select name, item_code from `tabSerial No` where item_code in (%s)"
		                     % ", ".join(["%s"] * len(items)), items)
	else:
		rows = frappe.db.sql("select name, item_code from `tabSerial No`")
	return {r[0]: r[1] for r in rows}


def sync_sims() -> dict:
	"""Every Lebara SIM into the `Lebara SIM` list, with its ERP vehicle. ~30,000 rows in about half a
	minute; only changed rows are written, and every row still listed is stamped `synced_at`."""
	started = now_datetime()
	erp = erp_vehicles()
	rule = serial_rule()
	serials = serial_index(rule)
	rows = []
	for e in iter_subscribers():
		row = sim_row(e)
		row.update(match_vehicle(row, erp))
		row.update(link_serial(row, serials, rule))
		rows.append(row)
	if not rows:
		raise LebaraError("Lebara returned no SIMs; nothing was changed.")
	res = store.upsert("Lebara SIM", "subscriber_id", rows)
	seconds = frappe.utils.time_diff_in_seconds(now_datetime(), started)
	linked = sum(1 for r in rows if r.get("serial_no"))
	res["linked_serials"] = linked
	res["note"] = "%s SIMs: %s new, %s changed, %s unchanged, %s with a Serial No (%ss)" % (
		res.get("total"), res.get("inserted"), res.get("updated"), res.get("unchanged"), linked, seconds)
	frappe.db.set_value(SETTINGS, SETTINGS, "lebara_sims_synced_at", now_datetime(), update_modified=False)
	frappe.db.set_value(SETTINGS, SETTINGS, "lebara_sims_count", cint(res.get("total")), update_modified=False)
	frappe.db.set_value(SETTINGS, SETTINGS, "lebara_sims_sync_note", res["note"], update_modified=False)
	frappe.db.commit()
	return res


def sync_invoice(year, month) -> dict:
	"""Store one month's invoice: the totals (`Lebara Invoice`) and every SIM's line
	(`Lebara Invoice Line`). Safe to repeat -- the same month is compared and only differences written."""
	got = fetch_invoice(year, month)
	if not got["ok"]:
		return {"ok": False, "period": got["period"], "note": got["note"]}
	period, now = got["period"], now_datetime()
	lines = []
	for ln in got["lines"]:
		lines.append(dict(ln, line_key="%s-%s" % (period, ln["msisdn"]), period=period))
	res = store.upsert("Lebara Invoice Line", "line_key", lines, scope={"period": period})
	total = sum(flt(ln["amount"]) for ln in got["lines"])
	head = dict(got["summary"], period=period, invoice_number=got["number"], invoice_date=got["date"],
	            billing_year=cint(year), billing_month=cint(month), lines_count=len(got["lines"]),
	            lines_total=total, synced_at=now)
	head.pop("number", None)
	store.upsert("Lebara Invoice", "period", [head])
	note = "%s: invoice %s, %s SIMs, %.2f SAR (%s new, %s changed)" % (
		period, got["number"], len(lines), total, res.get("inserted"), res.get("updated"))
	return {"ok": True, "period": period, "number": got["number"], "sims": len(lines), "lines_total": total,
	        "monthly_fee": got["summary"].get("monthly_fee"), "inserted": res.get("inserted"),
	        "updated": res.get("updated"), "note": note}


def sync_latest_invoice(today=None) -> dict:
	"""The invoice of the month before `today`, unless it is already stored with its lines. Meant for a
	daily job: most days it costs one database read and no request to Lebara."""
	year, month = previous_period(today)
	period = period_of(year, month)
	have = frappe.db.get_value("Lebara Invoice", period, ["lines_count", "invoice_number"], as_dict=True)
	if have and cint(have.lines_count):
		return {"ok": True, "period": period, "skipped": True, "note": "%s is already stored." % period}
	return sync_invoice(year, month)


def sync_transactions(page=PAGE) -> dict:
	"""Lebara's action history, incrementally. A transaction still Pending / Under Processing is read
	again until it settles, so a suspend that finished a minute after we first saw it is not left pending."""
	last = cint(frappe.db.sql("select max(transaction_id) from `tabLebara Transaction`")[0][0])
	# only recent ones: a transaction that never settles must not drag the window back to its own Id forever
	open_min = frappe.db.sql(
		"select min(transaction_id) from `tabLebara Transaction` "
		"where trans_status in (0, 1) and add_date > date_sub(now(), interval 7 day)")[0][0]
	since = min(last, cint(open_min) - 1) if open_min else last
	rows = [transaction_row(e) for e in iter_transactions(since, page)]
	if not rows:
		return {"ok": True, "total": 0, "inserted": 0, "updated": 0, "since_id": since,
		        "note": "no new transactions after %s" % since}
	res = store.upsert("Lebara Transaction", "transaction_id", rows, scope={"transaction_id": [">", since]})
	res.update(ok=True, since_id=since, note="%s transactions after %s (%s new, %s changed)" % (
		res.get("total"), since, res.get("inserted"), res.get("updated")))
	return res


SYNCS = ("sims", "invoice", "transactions")


def sync(what: str) -> dict:
	"""Refresh one mirror now: sims | invoice | transactions. Raises LebaraError on a problem."""
	steps = {"sims": sync_sims, "invoice": sync_latest_invoice, "transactions": sync_transactions}
	if what not in steps:
		raise LebaraError("Unknown sync: %s (use %s)" % (what, ", ".join(SYNCS)))
	return steps[what]()


def run_scheduled(what: str) -> dict:
	"""What the scheduler jobs call: skips quietly while Lebara is off or the session is not logged in,
	and returns the outcome instead of raising, so one bad day does not stop the next."""
	if not enabled() or _setting("lebara_session_status") != "Logged In":
		return {"ok": False, "skipped": True, "note": "Lebara is disabled or not logged in."}
	try:
		return sync(what)
	except LebaraError as e:
		frappe.log_error(title="Lebara sync %s" % what, message=str(e)[:1000])
		return {"ok": False, "note": str(e)[:300]}


# --------------------------------------------------------------------------
# the question the dashboard asks
# --------------------------------------------------------------------------


def _dead_sql(cfg: dict, as_of):
	"""(sql, params) for "the vehicle behind this SIM is gone", on Customer Vehicle's own columns, using
	the same rule as the Fleet Audit's waste (the fleet_audit_waste script): ERP status in a list, or a
	subscription that ran out more than N days before `as_of`."""
	parts, params = [], []
	statuses = [s for s in cfg.get("erp_statuses") or [] if s]
	if statuses:
		parts.append("cv.device_statues in (%s)" % ", ".join(["%s"] * len(statuses)))
		params += statuses
	if cfg.get("expired_days") is not None:
		parts.append("(cv.subscription_expiry_date is not null and cv.subscription_expiry_date < "
		             "date_sub(%s, interval %s day))")
		params += [as_of, cint(cfg["expired_days"])]
	if not parts:
		return "0", []
	return "case when " + " or ".join(parts) + " then 1 else 0 end", params


def loss_report(period: str, as_of=None) -> dict:
	"""What Lebara billed for `period` ('YYYY-MM') set against what the ERP knows about each SIM.

	Needs `sync_invoice` and `sync_sims` to have run. Every billed line (amount > 0) is grouped by the
	status Lebara billed it under, whether an ERP vehicle uses the SIM, and whether that vehicle is
	dead by the Fleet Audit's waste rule judged as of the end of the month. Amounts are SAR before VAT.

	-> {"period", "as_of", "billed": {sims, amount}, "unlinked": {...}, "dead_vehicle": {...},
	    "not_active": {...}, "groups": [{status, link, dead, sims, amount}], "invoice": {...}}
	"""
	if not re.fullmatch(r"\d{4}-\d{2}", str(period or "")):
		raise LebaraError("period must look like 2026-09.")
	year, month = int(period[:4]), int(period[5:])
	as_of = str(getdate(as_of) if as_of else get_last_day("%04d-%02d-01" % (year, month)))
	from app_apis import fleet_audit

	dead_sql, dead_params = _dead_sql(fleet_audit.waste_config(), as_of)
	rows = frappe.db.sql(
		"""
		select il.status,
		       case when s.msisdn is null then 'not_listed'
		            when ifnull(s.erp_vehicle, '') = '' then 'no_vehicle' else 'vehicle' end,
		       %s, count(*), sum(il.amount)
		from `tabLebara Invoice Line` il
		left join (select msisdn, max(erp_vehicle) erp_vehicle from `tabLebara SIM` group by msisdn) s
		       on s.msisdn = il.msisdn
		left join `tabCustomer Vehicle` cv on cv.name = s.erp_vehicle
		where il.period = %%s and il.amount > 0
		group by 1, 2, 3
		""" % dead_sql,
		[*dead_params, period],
	)
	found = [{"status": r[0] or "", "link": r[1], "dead": cint(r[2]), "sims": cint(r[3]), "amount": flt(r[4])}
	         for r in rows]

	def total(keep):
		sel = [g for g in found if keep(g)]
		return {"sims": sum(g["sims"] for g in sel), "amount": round(sum(g["amount"] for g in sel), 2)}

	inv = frappe.db.get_value("Lebara Invoice", period, ["invoice_number", "monthly_fee", "lines_total", "total",
	                                                     "synced_at"], as_dict=True)
	return {
		"period": period,
		"as_of": as_of,
		"billed": total(lambda g: True),
		"unlinked": total(lambda g: g["link"] != "vehicle"),
		"dead_vehicle": total(lambda g: g["dead"] and g["link"] == "vehicle"),
		"not_active": total(lambda g: g["status"] != "Active"),
		"groups": sorted(found, key=lambda g: -g["amount"]),
		"invoice": dict(inv) if inv else None,
	}
