# Copyright (c) 2026, osama and contributors
# For license information, please see license.txt

"""Unit tests for app_apis.core.lebara: the pure parsers and builders, the session handling with the
network mocked, and the incremental-sync window. No site and no request to Lebara.

    /home/frappe/frappe-bench/env/bin/python -m unittest app_apis.tests.test_lebara_core -v
"""

import io
import json
import os
import unittest
from unittest import mock

from app_apis.core import lebara, store

SCRIPTS = os.path.join(os.path.dirname(os.path.dirname(__file__)), "scripts")

INVOICE_HTML = """
<html><body><form id="BillingInvoice"><select name="BillingMonth"></select></form>
<div id="BillingInvoiceDetail"><div class="inv-details">
 <h1>INVOICE</h1>
 <table><tr><td>Invoice To:</td><td>M2M Intelligent Machines</td></tr>
 <tr><td>Invoice Number</td><td>158400000014987381</td></tr>
 <tr><td>Invoice Date</td><td>01 Oct 2026</td></tr>
 <tr><td>Billing Month</td><td>Sep 2026</td></tr></table>
</div></div><script>var x = "Invoice Number 999999999";</script></body></html>
"""


def workbook(numbers_rows, summary_rows=None, header=("Account ID", "MSISDN", "Tariff Plan", "Status", "Amount")):
	import openpyxl

	wb = openpyxl.Workbook()
	s = wb.active
	s.title = "Summary"
	for r in summary_rows or [("Number", "158400000014987381"), ("Previous Balance", 216200.92),
	                          ("Monthly Fee", 103397.09), ("Off-bundle Usage", 184.8), ("Adjustment", 0),
	                          ("Other service", 0), ("Total Payments", -120000), ("VAT 15%", 15537.28),
	                          ("Total SAR", 215320.08)]:
		s.append([None, *r])
	n = wb.create_sheet("Numbers")
	n.append(list(header))
	for r in numbers_rows:
		n.append(list(r))
	buf = io.BytesIO()
	wb.save(buf)
	return buf.getvalue()


class TestPureHelpers(unittest.TestCase):
	def test_dates_statuses_digits(self):
		self.assertEqual(lebara.lebara_dt("2026-09-16T10:20:57.000"), "2026-09-16 10:20:57")
		self.assertIsNone(lebara.lebara_dt(""))
		self.assertEqual(lebara.status_text(2), "Active")
		self.assertEqual(lebara.status_text(9), "Deactivated")
		self.assertEqual(lebara.status_text(77), "")
		self.assertEqual(lebara.digits(" 830-033 "), "830033")

	def test_invoice_status_uses_the_sim_list_vocabulary(self):
		self.assertEqual(lebara.invoice_status("Deactivation"), "Deactivated")
		self.assertEqual(lebara.invoice_status(" active "), "Active")
		self.assertEqual(lebara.invoice_status("Suspend"), "Suspend")
		self.assertEqual(lebara.invoice_status("Something New"), "Something New")  # never silently renamed

	def test_previous_period(self):
		self.assertEqual(lebara.previous_period("2026-10-10"), (2026, 9))
		self.assertEqual(lebara.previous_period("2026-01-31"), (2025, 12))
		self.assertEqual(lebara.period_of(2026, 9), "2026-09")

	def test_list_body(self):
		b = lebara.list_body(skip=10, take=999999, status=4)
		self.assertEqual(b["Take"], lebara.PAGE)
		self.assertEqual(b["EqualityFilter"], {"SubscriberStatusM2M1": 4})
		self.assertEqual(b["Skip"], 10)
		self.assertNotIn("Criteria", b)
		q = lebara.list_body(q="abc 8300-33")
		self.assertEqual(q["Criteria"][0][0][2], "%830033%")
		self.assertIn("ICCID", json.dumps(q["Criteria"]))
		self.assertNotIn("EqualityFilter", lebara.list_body(status=""))

	def test_sim_row_and_match(self):
		row = lebara.sim_row({"Id": 7, "Msisdn": 83003, "ICCID": "8996", "SubscriberStatusM2M1": 4, "IMEI": "86",
		                      "LastDayUsage": 1.5, "IMEI_Date": "2026-01-02T03:04:05.000", "Apn": "lebara"})
		self.assertEqual((row["subscriber_id"], row["msisdn"], row["status"], row["last_day_mb"]), (7, "83003", "Suspend", 1.5))
		self.assertEqual(row["imei_date"], "2026-01-02 03:04:05")
		erp = {"iccid": {"8996": {"erp_vehicle": "V1", "erp_customer": "C", "erp_plate": "P", "erp_imei": "86"}},
		       "imei": {"99": {"erp_vehicle": "V2", "erp_customer": "", "erp_plate": "Q", "erp_imei": "99"}}}
		self.assertEqual(lebara.match_vehicle(row, erp)["match_by"], "ICCID")
		self.assertEqual(lebara.match_vehicle({"iccid": "none", "imei": "99"}, erp)["erp_vehicle"], "V2")
		self.assertEqual(lebara.match_vehicle({"iccid": "none", "imei": "99"}, erp)["match_by"], "IMEI")
		none = lebara.match_vehicle({"iccid": "x", "imei": ""}, erp)
		self.assertEqual((none["erp_vehicle"], none["match_by"]), (None, ""))

	def test_transaction_row(self):
		t = lebara.transaction_row({"Id": 5, "SubscriberId": 9, "TransTypeM2M1": 5, "Msisdn": "830", "TransStatus": 2,
		                            "TransResult": "SIM is suspended", "AddDate": "2026-10-07T08:51:38.000", "AddBy": 586})
		self.assertEqual((t["transaction_id"], t["trans_type_text"], t["trans_status_text"]), (5, "Suspend SIM", "Success"))
		self.assertEqual(t["add_date"], "2026-10-07 08:51:38")
		blank = lebara.transaction_row({"Id": 1})
		self.assertEqual((blank["trans_type_text"], blank["trans_status_text"]), ("", ""))

	def test_logged_out_detection(self):
		lo = lebara.looks_logged_out
		self.assertTrue(lo(None))
		self.assertTrue(lo({"url": "https://b2b.lebara.sa/Account/Login?ReturnUrl=x", "status": 200}))
		self.assertTrue(lo({"status": 403, "url": "u"}))
		self.assertTrue(lo({"status": 400, "url": "u", "json": {"Error": {"Code": "NotLoggedIn"}}}))
		self.assertTrue(lo({"status": 200, "url": "u", "json": None, "text": "<a href='/Account/Login'>"}))
		self.assertFalse(lo({"status": 200, "url": "u", "json": {"Entities": []}}))
		self.assertFalse(lo({"status": 500, "url": "u", "json": {"Error": {"Code": "Exception"}}}))


class TestInvoice(unittest.TestCase):
	def test_page(self):
		p = lebara.parse_invoice_page(INVOICE_HTML)
		self.assertEqual(p, {"number": "158400000014987381", "date": "2026-10-01", "billing_month": "Sep 2026"})
		self.assertIsNone(lebara.parse_invoice_page("<html>Select bill cycle month and year</html>"))
		self.assertIsNone(lebara.parse_invoice_page(""))

	def test_workbook(self):
		data = workbook([(None, "830033297398", "M2M OFFER LEBARA", "Active", 5),
		                 (None, "830033297403", "M2M OFFER LEBARA", "Deactivation", 2.5),
		                 (None, "830033297413", "M2M OFFER LEBARA", "Idle", 0),
		                 (None, None, None, None, None)])
		got = lebara.parse_invoice_workbook(data)
		self.assertEqual([l["msisdn"] for l in got["lines"]], ["830033297398", "830033297403", "830033297413"])
		self.assertEqual([l["status"] for l in got["lines"]], ["Active", "Deactivated", "Idle"])
		self.assertEqual(sum(l["amount"] for l in got["lines"]), 7.5)
		s = got["summary"]
		self.assertEqual((s["monthly_fee"], s["off_bundle_usage"], s["total_payments"]), (103397.09, 184.8, -120000.0))
		self.assertEqual(s["number"], "158400000014987381")
		self.assertAlmostEqual(s["total"], 215320.08)

	def test_columns_are_found_by_name_not_position(self):
		data = workbook([("Sep", "Active", 5, "830033297398", "PLAN")],
		                header=("Period", "Status", "Amount", "MSISDN", "Tariff Plan"))
		line = lebara.parse_invoice_workbook(data)["lines"][0]
		self.assertEqual((line["msisdn"], line["status"], line["amount"], line["tariff_plan"]),
		                 ("830033297398", "Active", 5.0, "PLAN"))

	def test_not_an_invoice(self):
		import openpyxl

		wb = openpyxl.Workbook()
		buf = io.BytesIO()
		wb.save(buf)
		with self.assertRaises(lebara.LebaraError):
			lebara.parse_invoice_workbook(buf.getvalue())
		with self.assertRaises(lebara.LebaraError):
			lebara.parse_invoice_workbook(b"not a zip")
		with self.assertRaises(lebara.LebaraError):
			lebara.parse_invoice_workbook(workbook([], header=("a", "b")))


class TestSession(unittest.TestCase):
	OK = {"status": 200, "url": "u", "json": {"Entities": [{"Id": 1}], "TotalCount": 1}}
	OUT = {"status": 403, "url": "u", "json": None, "text": ""}

	def setUp(self):
		self.patches = [
			mock.patch.object(lebara, "refresh_csrf"),
			mock.patch.object(lebara, "set_state"),
			mock.patch.object(lebara, "_notify_expired"),
			mock.patch.object(lebara, "_setting", return_value="Logged In"),
			mock.patch.object(lebara, "frappe", mock.MagicMock()),
		]
		self.m = [p.start() for p in self.patches]
		self.addCleanup(lambda: [p.stop() for p in self.patches])

	def test_one_transient_logged_out_answer_does_not_expire_the_session(self):
		with mock.patch.object(lebara, "_post_json", side_effect=[self.OUT, self.OK]) as post:
			self.assertEqual(lebara.service("/x", {})["TotalCount"], 1)
		self.assertEqual(post.call_count, 2)
		lebara.set_state.assert_not_called()

	def test_two_logged_out_answers_expire_it_once_and_alert_once(self):
		with mock.patch.object(lebara, "_post_json", return_value=self.OUT):
			with self.assertRaises(lebara.SessionExpired):
				lebara.service("/x", {})
		lebara.set_state.assert_called_once()
		self.assertEqual(lebara.set_state.call_args[0][0], "Expired")
		lebara._notify_expired.assert_called_once()

	def test_already_expired_is_not_alerted_again(self):
		with mock.patch.object(lebara, "_setting", return_value="Expired"), \
				mock.patch.object(lebara, "_post_json", return_value=self.OUT):
			with self.assertRaises(lebara.SessionExpired):
				lebara.service("/x", {})
		lebara.set_state.assert_not_called()
		lebara._notify_expired.assert_not_called()

	def test_stale_csrf_gets_one_page_load_and_one_retry(self):
		err = {"status": 500, "url": "u", "json": {"Error": {"Code": "Exception"}}}
		with mock.patch.object(lebara, "_post_json", side_effect=[err, self.OK]):
			self.assertTrue(lebara.service("/x", {}))
		lebara.refresh_csrf.assert_called_once()

	def test_network_error_and_service_error_and_no_json(self):
		with mock.patch.object(lebara, "_post_json", return_value={"error": "timeout", "status": 0}):
			with self.assertRaises(lebara.LebaraError) as c:
				lebara.service("/x", {})
		self.assertIn("unreachable", str(c.exception))
		bad = {"status": 400, "url": "u", "json": {"Error": {"Code": "AccessDenied", "Message": "No."}}}
		with mock.patch.object(lebara, "_post_json", return_value=bad):
			with self.assertRaises(lebara.LebaraError) as c:
				lebara.service("/x", {})
		self.assertIn("No.", str(c.exception))
		with mock.patch.object(lebara, "_post_json", return_value={"status": 200, "url": "u", "json": None, "text": "x"}):
			with self.assertRaises(lebara.LebaraError):
				lebara.service("/x", {})

	def test_paging_stops_at_total_and_on_an_empty_page(self):
		pages = [{"Entities": [{"Id": i} for i in range(2)], "TotalCount": 5},
		         {"Entities": [{"Id": i} for i in range(2, 4)], "TotalCount": 5},
		         {"Entities": [{"Id": 4}], "TotalCount": 5}]
		with mock.patch.object(lebara, "service", side_effect=pages) as svc:
			got = list(lebara._paged("/p", {"Sort": ["Id"]}, page=2))
		self.assertEqual([g["Id"] for g in got], [0, 1, 2, 3, 4])
		self.assertEqual([c[0][1]["Skip"] for c in svc.call_args_list], [0, 2, 4])
		with mock.patch.object(lebara, "service", return_value={"Entities": [], "TotalCount": 99}):
			self.assertEqual(list(lebara._paged("/p", {}, page=2)), [])


class TestSyncWindows(unittest.TestCase):
	def _run(self, last, open_min):
		fr = mock.MagicMock()
		fr.db.sql.side_effect = [[[last]], [[open_min]]]
		with mock.patch.object(lebara, "frappe", fr), \
				mock.patch.object(lebara, "iter_transactions", return_value=iter([{"Id": 1}])) as it, \
				mock.patch.object(lebara.store, "upsert", return_value={"total": 1, "inserted": 1, "updated": 0}) as up:
			lebara.sync_transactions()
		return it.call_args[0][0], up.call_args[1]["scope"]

	def test_normal_run_reads_after_the_newest_stored_id(self):
		since, scope = self._run(1000, None)
		self.assertEqual(since, 1000)
		self.assertEqual(scope, {"transaction_id": [">", 1000]})

	def test_a_pending_transaction_pulls_the_window_back_to_it(self):
		self.assertEqual(self._run(1000, 940)[0], 939)

	def test_empty_table_reads_everything(self):
		self.assertEqual(self._run(None, None)[0], 0)

	def test_nothing_new_writes_nothing(self):
		fr = mock.MagicMock()
		fr.db.sql.side_effect = [[[5]], [[None]]]
		with mock.patch.object(lebara, "frappe", fr), \
				mock.patch.object(lebara, "iter_transactions", return_value=iter([])), \
				mock.patch.object(lebara.store, "upsert") as up:
			out = lebara.sync_transactions()
		self.assertTrue(out["ok"])
		up.assert_not_called()

	def test_latest_invoice_is_skipped_when_stored(self):
		fr = mock.MagicMock()
		fr.db.get_value.return_value = mock.Mock(lines_count=26000, invoice_number="1")
		with mock.patch.object(lebara, "frappe", fr), mock.patch.object(lebara, "sync_invoice") as si:
			out = lebara.sync_latest_invoice("2026-10-10")
		self.assertTrue(out["skipped"])
		si.assert_not_called()
		fr = mock.MagicMock()
		fr.db.get_value.return_value = None
		with mock.patch.object(lebara, "frappe", fr), \
				mock.patch.object(lebara, "sync_invoice", return_value={"ok": True}) as si:
			lebara.sync_latest_invoice("2026-10-10")
		si.assert_called_once_with(2026, 9)

	def test_scheduled_run_skips_quietly_and_never_raises(self):
		with mock.patch.object(lebara, "enabled", return_value=False):
			self.assertTrue(lebara.run_scheduled("sims")["skipped"])
		with mock.patch.object(lebara, "enabled", return_value=True), \
				mock.patch.object(lebara, "_setting", return_value="Logged In"), \
				mock.patch.object(lebara, "sync", side_effect=lebara.LebaraError("boom")), \
				mock.patch.object(lebara, "frappe", mock.MagicMock()) as fr:
			log = fr.log_error
			out = lebara.run_scheduled("sims")
		self.assertFalse(out["ok"])
		log.assert_called_once()
		with self.assertRaises(lebara.LebaraError):
			lebara.sync("nope")

	def test_dead_vehicle_sql_binds_everything(self):
		sql, params = lebara._dead_sql({"erp_statuses": ["Deleted", "x'; --"], "expired_days": 30}, "2026-09-30")
		self.assertNotIn("--", sql)
		self.assertEqual(params, ["Deleted", "x'; --", "2026-09-30", 30])
		self.assertEqual(lebara._dead_sql({"erp_statuses": [], "expired_days": None}, "2026-09-30"), ("0", []))


class TestStoreSeen(unittest.TestCase):
	def test_mark_seen_chunks_and_never_touches_modified(self):
		with mock.patch.object(store, "frappe", mock.MagicMock()) as fr:
			sql = fr.db.sql
			store.mark_seen("Lebara SIM", [str(i) for i in range(4500)], "T")
		self.assertEqual(sql.call_count, 3)  # 2000 + 2000 + 500
		text = sql.call_args_list[0][0][0]
		self.assertIn("`synced_at` = %s", text)
		self.assertNotIn("modified", text)
		self.assertEqual(len(sql.call_args_list[0][0][1]), 2001)


class TestSerialLink(unittest.TestCase):
	INDEX = {"8996606099014667123": "Lebara SIM", "8996606099000000001": "sim_stc", "830033": "Lebara SIM"}

	def test_default_rule_links_by_iccid(self):
		rule = lebara.clean_serial_rule(None)
		self.assertEqual(rule, {"match_fields": ["iccid"], "item_codes": []})
		got = lebara.link_serial({"iccid": "8996606099014667123", "msisdn": "x"}, self.INDEX, rule)
		self.assertEqual(got, {"serial_no": "8996606099014667123", "serial_item": "Lebara SIM"})
		none = lebara.link_serial({"iccid": "nope"}, self.INDEX, rule)
		self.assertEqual(none, {"serial_no": None, "serial_item": None})

	def test_fields_are_tried_in_the_scripts_order(self):
		rule = lebara.clean_serial_rule({"match_fields": ["msisdn", "iccid"]})
		row = {"iccid": "8996606099014667123", "msisdn": "830033"}
		self.assertEqual(lebara.link_serial(row, self.INDEX, rule)["serial_no"], "830033")
		row["msisdn"] = "unknown"
		self.assertEqual(lebara.link_serial(row, self.INDEX, rule)["serial_no"], "8996606099014667123")

	def test_a_bad_script_answer_keeps_the_default(self):
		d = lebara.clean_serial_rule(None)
		for bad in ("x", 5, [], {"match_fields": ["colour"]}, {"match_fields": []}, {"match_fields": ["iccid", "bogus"]}):
			self.assertEqual(lebara.clean_serial_rule(bad), d, bad)
		self.assertEqual(lebara.clean_serial_rule({"match_fields": "MSISDN, iccid"})["match_fields"], ["msisdn", "iccid"])
		self.assertEqual(lebara.clean_serial_rule({"item_codes": "Lebara SIM, sim_stc"})["item_codes"], ["Lebara SIM", "sim_stc"])

	def test_defaults_are_not_shared_between_calls(self):
		lebara.clean_serial_rule(None)["match_fields"].append("imsi")
		self.assertEqual(lebara.SERIAL_DEFAULTS["match_fields"], ["iccid"])

	def test_the_index_query_is_restricted_to_the_rules_items(self):
		fr = mock.MagicMock()
		fr.db.sql.return_value = [("A", "Lebara SIM")]
		with mock.patch.object(lebara, "frappe", fr):
			self.assertEqual(lebara.serial_index({"item_codes": ["Lebara SIM", "sim_stc"]}), {"A": "Lebara SIM"})
		sql, params = fr.db.sql.call_args[0]
		self.assertIn("item_code in (%s, %s)", sql)
		self.assertEqual(params, ["Lebara SIM", "sim_stc"])
		with mock.patch.object(lebara, "frappe", fr):
			lebara.serial_index({"item_codes": []})
		self.assertNotIn("where", fr.db.sql.call_args[0][0])


class TestShippedScripts(unittest.TestCase):
	def _meta(self, name):
		with open(os.path.join(SCRIPTS, "server_scripts.json")) as f:
			return next(m for m in json.load(f) if m["name"] == name)

	def test_lebara_scripts_are_managed_and_compile_in_the_sandbox(self):
		from frappe.utils.safe_exec import FrappeTransformer
		from RestrictedPython import compile_restricted

		for name in ("Lebara API", "Lebara SIM Sync", "Lebara Keepalive", "Lebara History Sync", "lebara_serial_link"):
			meta = self._meta(name)
			self.assertTrue(meta["managed"], name)
			with open(os.path.join(SCRIPTS, "server", meta["file"])) as f:
				compile_restricted(f.read(), name, "exec", policy=FrappeTransformer)
		self.assertEqual(self._meta("Lebara SIM Sync")["carry"], ["QUIET_FROM", "QUIET_TO"])
		self.assertEqual(self._meta("Lebara History Sync")["carry"], ["INVOICE_HOUR"])
		self.assertEqual(self._meta("lebara_serial_link")["carry"], ["MATCH_FIELDS", "ITEM_CODES"])

	def test_the_sandbox_sync_delegates_to_the_core(self):
		with open(os.path.join(SCRIPTS, "server", "lebara_api.py")) as f:
			src = f.read()
		self.assertIn('frappe.call("app_apis.core.api.lebara_sync", what="sims")', src)
		self.assertNotIn("def erp_vehicles", src)

	def test_new_doctypes_are_named_by_their_key_so_upsert_works(self):
		base = os.path.join(os.path.dirname(os.path.dirname(__file__)), "app_apis", "doctype")
		for folder, key in (("lebara_invoice", "period"), ("lebara_invoice_line", "line_key"),
		                    ("lebara_transaction", "transaction_id"), ("lebara_sim", "subscriber_id")):
			with open(os.path.join(base, folder, folder + ".json")) as f:
				d = json.load(f)
			self.assertEqual(d["autoname"], "field:" + key, folder)
			names = {x["fieldname"] for x in d["fields"]}
			self.assertIn(key, names)
			self.assertIn("synced_at", names, folder)


if __name__ == "__main__":
	unittest.main()
